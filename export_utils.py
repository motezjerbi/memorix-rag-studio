"""
export_utils.py
---------------
Génère des exports PDF et Word (.docx) prêts à l'impression à partir des synthèses
et fiches mémo produites par rag_core.py, avec prise en charge du formatage de code
et mise en valeur des formules.

Dépendances requises :
    pip install reportlab python-docx
"""

import io
import re

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable, Preformatted

from docx import Document as DocxDocument
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def build_export_blocks(title: str, result: dict):
    """
    Transforme un résultat structuré en flux de blocs typés :
    (type, contenu) avec type ∈ {"title", "heading", "paragraph", "code", "rule"}.
    """
    blocks = [("title", title)]

    if result.get("synthesis"):
        blocks.append(("heading", "Synthèse globale & Fondements"))
        for line in result["synthesis"].split("\n"):
            line_s = line.strip()
            if line_s:
                blocks.append(("paragraph", line_s))
        blocks.append(("rule", ""))

    for s in result.get("sections", []):
        if s.get("label"):
            blocks.append(("heading", s["label"]))

        in_code_block = False
        code_lines = []

        for line in s["text"].split("\n"):
            raw_line = line.rstrip()
            if raw_line.strip().startswith("```"):
                if in_code_block:
                    blocks.append(("code", "\n".join(code_lines)))
                    code_lines = []
                    in_code_block = False
                else:
                    in_code_block = True
                continue

            if in_code_block:
                code_lines.append(raw_line)
            elif raw_line.strip():
                blocks.append(("paragraph", raw_line.strip()))

        if code_lines:
            blocks.append(("code", "\n".join(code_lines)))

        blocks.append(("rule", ""))

    return blocks


def _escape_xml(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _markdown_to_reportlab(text: str) -> str:
    """Convertit le balisage gras, italique et code inline pour ReportLab."""
    text = _escape_xml(text)
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"\*(.+?)\*", r"<i>\1</i>", text)
    text = re.sub(r"`(.+?)`", r'<font face="Courier" color="#7C5CFC">\1</font>', text)
    return text


def export_to_pdf_bytes(title: str, result: dict) -> bytes:
    """Génère un rapport PDF élégant en mémoire et retourne les octets."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleDS",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor="#1E1B4B",
        alignment=TA_LEFT,
        spaceAfter=14,
    )
    heading_style = ParagraphStyle(
        "HeadingDS",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=17,
        textColor="#4338CA",
        spaceBefore=14,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "BodyDS",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14.5,
        textColor="#1F2937",
        spaceAfter=5,
    )
    code_style = ParagraphStyle(
        "CodeDS",
        parent=styles["Normal"],
        fontName="Courier",
        fontSize=8.5,
        leading=11,
        textColor="#111827",
        backColor="#F3F4F6",
        borderPadding=6,
        spaceAfter=6,
    )

    story = []
    for kind, text in build_export_blocks(title, result):
        if kind == "title":
            story.append(Paragraph(_escape_xml(text), title_style))
            story.append(HRFlowable(width="100%", thickness=2, color="#4F46E5", spaceAfter=14))
        elif kind == "heading":
            story.append(Paragraph(_escape_xml(text), heading_style))
        elif kind == "rule":
            story.append(Spacer(1, 3))
            story.append(HRFlowable(width="100%", thickness=0.5, color="#E5E7EB", spaceAfter=8))
        elif kind == "code":
            story.append(Preformatted(text, code_style))
        else:
            story.append(Paragraph(_markdown_to_reportlab(text), body_style))

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def export_to_docx_bytes(title: str, result: dict) -> bytes:
    """Génère un document Word (.docx) structuré en mémoire vive."""
    doc = DocxDocument()

    h0 = doc.add_heading(title, level=0)
    h0.paragraph_format.space_after = Pt(14)

    for kind, text in build_export_blocks(title, result):
        if kind == "title":
            continue
        elif kind == "heading":
            h = doc.add_heading(text, level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
        elif kind == "rule":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("―" * 30)
            run.font.color.rgb = RGBColor(209, 213, 219)
        elif kind == "code":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(text)
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(31, 41, 55)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            for part in re.split(r"(\*\*.+?\*\*|`.+?`)", text):
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                elif part.startswith("`") and part.endswith("`"):
                    r = p.add_run(part[1:-1])
                    r.font.name = "Courier New"
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(124, 92, 252)
                elif part:
                    p.add_run(part)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()